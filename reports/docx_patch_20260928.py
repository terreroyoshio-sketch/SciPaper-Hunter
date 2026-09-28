"""One-off patch for 制图提示词.docx (2026-09-28 rule correction).

Creates a patched COPY (original untouched): 制图提示词.patched_20260928.docx
Logs every change to reports/docx_patch_log_20260928.md.
Reproduce: python reports/docx_patch_20260928.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

SRC = Path(r"C:\Users\张涵\Desktop\制图提示词.docx")
DST = Path(r"C:\Users\张涵\Desktop\制图提示词.patched_20260928.docx")
LOG = Path(__file__).resolve().parent / "docx_patch_log_20260928.md"

REPLACEMENTS = [
    ("* research-task-router", "* scientific-figure-router"),
    ("* diagram-workflow", "* figure-quality-gate"),
    ("* diagram-quality-engineering", "* figurelab-publish"),
    ("* diagram-visual-output-guard", "* academic-figure-skill"),
    ("* cognitive-illustration", "* scientific-visualization"),
    ("* scientific-workbench-agent", "* nature-figure-style"),
    ("* final-delivery-audit", "* technical-route-diagram"),
    ("* paper-workflow", "* figure-preflight-checker"),
    ("## 绝对禁止传统：", "## 禁止对象（2026-09-28 修正）：无语义模板化流程图"),
    ("* 大量方框 + 箭头", "* 无语义、模板化的方框堆叠（大量同尺寸方框 + 箭头直连）"),
    ("# 十八、流程图必须采用“几何叙事”", "# 十八、流程图规则（2026-09-28 修正）：分场景处理"),
    ("视觉上不得画成传统流程图。", "视觉上不得画成“无语义模板化”的默认流程图；本类（数学建模算法）图应优先几何叙事。"),
    ("19. 流程图不是方框+箭头；", "19. 流程图不得是“无语义模板化方框堆叠”（默认 Mermaid/SmartArt 风、等尺寸节点直连）；允许承担信息组织功能的语义容器 + 箭头；"),
    ("禁止方框堆叠。", "禁止无语义、模板化的方框堆叠（允许承担信息组织功能的语义容器 + 箭头）。"),
]

INSERT_AFTER = {
    "* AI 信息图风格": [
        "（2026-09-28 修正）",
        "**禁止**无语义、模板化的方框堆叠和默认流程图风格（默认 Mermaid 风、PPT SmartArt、彩色圆角卡片堆、全节点同尺寸同箭头、纯文字节点占据全部视觉主体、装饰性图标）。",
        "**允许**矩形、圆角矩形、分区面板、箭头、连接线、虚线区域、公式、地图、数据小图、几何示意——但每个容器必须承担明确的信息组织功能。",
        "技术路线图与方法框架图优先采用“数据层—方法层—分析层—结果层”或与论文实际方法一致的层级；重要步骤应结合真实小图、公式、地图、几何对象或结果缩略图，而不是全部用文字节点表达。",
    ],
    "# 十八、流程图规则（2026-09-28 修正）：分场景处理": [
        "几何叙事（用点 / 射线 / 区域 / 轨迹等与数学模型对应的几何对象表达算法阶段，视觉信息 ≥60% 来自几何关系）仍是**数学建模类算法流程图**的优先路径（本节后续几何对象清单与示例继续适用）。",
        "但**方法框架图、技术路线图、一般工程与数据流程**允许使用语义容器 + 箭头，前提是语义明确、不得是模板化堆叠。",
    ],
}

HEADER_NOTE = (
    "【2026-09-28 修订副本】本副本修正了“方框+箭头 / 几何叙事”绝对规则"
    "（§十七 / §十八 / §五十二 / §五十四），并将技能列表同步为实际已装技能；"
    "生成方式为脚本补丁，原件未改动。变更明细见 docx_patch_log_20260928.md。"
)


def ptext(p) -> str:
    return "".join(r.text for r in p.runs)


def set_text(p, new: str) -> None:
    if p.runs:
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(new)


def insert_paragraph_after(paragraph, text: str):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.add_run(text)
    return new_para


def main() -> None:
    doc = Document(str(SRC))
    log: list[tuple[str, str, str]] = []
    applied = {old: 0 for old, _ in REPLACEMENTS}
    inserted: set[str] = set()

    doc.paragraphs[0].insert_paragraph_before(HEADER_NOTE)
    log.append(("INSERT 文档头修订说明", "（无）", HEADER_NOTE))

    for i, p in enumerate(doc.paragraphs):
        t = ptext(p)
        if not t:
            continue
        for old, new in REPLACEMENTS:
            if t == old:
                set_text(p, new)
                applied[old] += 1
                log.append((f"REPLACE [段 {i}]", old, new))
                break
        for key, lines in INSERT_AFTER.items():
            if ptext(p) == key and key not in inserted:
                anchor = p
                for line in lines:
                    anchor = insert_paragraph_after(anchor, line)
                inserted.add(key)
                log.append((f"INSERT-AFTER [段 {i}] 锚点：{key[:26]}…", "（无）", " ‖ ".join(lines)))

    missing = [old for old, c in applied.items() if c == 0]
    doc.save(str(DST))

    lines_out = ["# docx 补丁日志（2026-09-28）", "", f"- 源文件：`{SRC}`", f"- 输出副本：`{DST}`", f"- 变更条数：{len(log)}", ""]
    for tag, old, new in log:
        lines_out += [f"## {tag}", f"- 旧：{old}", f"- 新：{new}", ""]
    if missing:
        lines_out += ["## 未命中（需人工核对）"] + [f"- {m}" for m in missing]
    LOG.write_text("\n".join(lines_out), encoding="utf-8")

    print(f"saved: {DST}")
    print(f"log  : {LOG}")
    print(f"edits: {len(log)}; missing: {len(missing)}")
    for m in missing:
        print("MISS:", m)


if __name__ == "__main__":
    main()
