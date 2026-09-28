"""Publication three-line table builder (python-docx, 2026-09-28).

Implements the reviewer spec:
- THREE rules only (top / header separator / bottom); all other cell borders nil.
- Table-level CENTER (w:jc via Table.alignment) is set - this is NOT the same as
  paragraph alignment inside cells (python-docx docs: alignment positions the
  table between the margins).
- Column widths are explicitly computed from content (CJK full-width / Latin
  half-width estimate), table layout = FIXED, autofit OFF (no Word drift).
- Per-column paragraph alignment by semantics; horizontal centering of the whole
  table is separate from cell text alignment.
- Vertical centering in cells; header row repeats across pages.
- Caption above the table, left-aligned at the text margin.
"""

from __future__ import annotations

from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

_AMAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


def text_width_mm(doc) -> float:
    sec = doc.sections[0]
    emu = int(sec.page_width) - int(sec.left_margin) - int(sec.right_margin)
    return emu / 36000.0  # 36000 EMU per mm


def style_document(doc, latin: str = "Times New Roman", east_asia: str = "SimSun", size_pt: float = 10.5) -> None:
    """Set Normal font (Latin + EastAsia) without touching user templates."""
    style = doc.styles["Normal"]
    style.font.name = latin
    style.font.size = Pt(size_pt)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)


def _cell_borders(cell, top=None, bottom=None, left=None, right=None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_b = tc_pr.find(qn("w:tcBorders"))
    if tc_b is None:
        tc_b = OxmlElement("w:tcBorders")
        tc_pr.append(tc_b)
    for edge, spec in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = tc_b.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tc_b.append(el)
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            sz, val = spec
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(sz))  # eighths of a point
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")


def _table_fixed_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)


def _cell_margins(table, top=40, bottom=40, left=80, right=80) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for name, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def natural_width_mm(headers, rows, mm_per_unit: float = 3.7, cell_pad_mm: float = 4.6, slack_mm: float = 3.0) -> float:
    """Content-based natural table width: short tables must NOT be forced to the
    full text width (reviewer rule); long ones may approach it."""

    def u(s) -> float:
        s = str(s)
        full = sum(1 for c in s if ord(c) > 0x2E7F)
        return full * 1.0 + (len(s) - full) * 0.55

    total = 0.0
    for c in range(len(headers)):
        mx = u(headers[c])
        for r in rows:
            mx = max(mx, max((u(x) for x in str(r[c]).split("\n")), default=0.0))
        total += max(mx, 3.0) * mm_per_unit + cell_pad_mm
    return total + slack_mm


def estimate_col_widths(headers, rows, total_mm: float, min_mm: float = 10.0):
    def w(s) -> float:
        s = str(s)
        full = sum(1 for c in s if ord(c) > 0x2E7F)
        other = len(s) - full
        return full * 1.0 + other * 0.55

    n = len(headers)
    weights = []
    for c in range(n):
        mx = w(headers[c])
        for r in rows:
            mx = max(mx, max((w(x) for x in str(r[c]).split("\n")), default=0.0))
        weights.append(max(mx, 3.0))
    unit = total_mm / sum(weights)
    widths = [max(min_mm, unit * x) for x in weights]
    scale = total_mm / sum(widths)
    return [round(x * scale, 1) for x in widths]


def add_three_line_table(doc, headers, rows, col_align=None, total_width_mm=None, caption=None):
    text_w = text_width_mm(doc)
    nat = natural_width_mm(headers, rows)
    tw = min(total_width_mm or nat, text_w)
    tw = max(tw, min(60.0, text_w))
    widths = estimate_col_widths(headers, rows, tw)
    n = len(headers)
    aligns = col_align or (["center"] * n)

    if caption:
        p = doc.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(4)

    table = doc.add_table(rows=1 + len(rows), cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER  # whole-table placement between margins
    table.autofit = False
    _table_fixed_layout(table)
    _cell_margins(table)

    for j, wmm in enumerate(widths):
        table.columns[j].width = Mm(wmm)

    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.width = Mm(widths[j])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        para.alignment = _AMAP[aligns[j]]
        run = para.add_run(str(h))
        run.bold = True
        _cell_borders(cell, top=(12, "single"), bottom=(8, "single"))

    for i, r in enumerate(rows, start=1):
        last = i == len(rows)
        for j, val in enumerate(r):
            cell = table.cell(i, j)
            cell.width = Mm(widths[j])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            para = cell.paragraphs[0]
            para.alignment = _AMAP[aligns[j]]
            para.add_run(str(val))
            _cell_borders(cell, bottom=(12, "single") if last else None)

    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))
    return table
