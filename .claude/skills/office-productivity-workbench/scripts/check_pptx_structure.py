#!/usr/bin/env python3
"""Check PPTX structure: slides, text, images, tables.
Usage: check_pptx_structure.py <input_pptx> [output_dir]
"""

import sys
import os
from datetime import datetime

try:
    from pptx import Presentation
except ImportError:
    print("[ERROR] python-pptx not installed. Run: pip install python-pptx")
    sys.exit(1)


def check_structure(input_path, output_dir=None):
    prs = Presentation(input_path)

    report_lines = [
        f"# PPTX Structure Check",
        f"",
        f"**File:** {os.path.basename(input_path)}",
        f"**Slides:** {len(prs.slides)}",
        f"**Slide width:** {prs.slide_width}",
        f"**Slide height:** {prs.slide_height}",
        f"**Analysed:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"",
    ]

    for i, slide in enumerate(prs.slides, 1):
        texts = []
        has_image = False
        has_table = False

        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    texts.append(text[:100])
            if hasattr(shape, "image"):
                has_image = True
            if shape.has_table:
                has_table = True

        report_lines.append(f"## Slide {i}")
        report_lines.append(f"- Shapes: {len(slide.shapes)}")
        report_lines.append(f"- Images: {'Yes' if has_image else 'No'}")
        report_lines.append(f"- Tables: {'Yes' if has_table else 'No'}")
        if texts:
            # Show first text as title candidate
            report_lines.append(f"- First text: {texts[0]}")
        report_lines.append("")

    if output_dir:
        os.makedirs(output_dir, exist_ok=True)
        out_path = os.path.join(output_dir, "pptx_structure_check.md")
        with open(out_path, "w", encoding="utf-8") as f:
            f.write("\n".join(report_lines))
        print(f"[OK] Check report written to {out_path}")
    else:
        print("\n".join(report_lines))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: check_pptx_structure.py <input_pptx> [output_dir]")
        sys.exit(1)
    check_structure(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
