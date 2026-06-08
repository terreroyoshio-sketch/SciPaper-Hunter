#!/usr/bin/env python3
"""Check DOCX structure: headings, paragraphs, tables.
Usage: check_docx_structure.py <input_docx> [output_dir]
"""

import sys
import os
from datetime import datetime

try:
    from docx import Document
except ImportError:
    print("[ERROR] python-docx not installed. Run: pip install python-docx")
    sys.exit(1)


def check_structure(input_path, output_dir=None):
    doc = Document(input_path)

    report_lines = [
        f"# DOCX Structure Check",
        f"",
        f"**File:** {os.path.basename(input_path)}",
        f"**Analysed:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"",
    ]

    # Count elements
    paras = len(doc.paragraphs)
    tables = len(doc.tables)
    headings = sum(1 for p in doc.paragraphs if p.style.name.startswith("Heading"))

    report_lines.append(f"**Paragraphs:** {paras}")
    report_lines.append(f"**Headings:** {headings}")
    report_lines.append(f"**Tables:** {tables}")
    report_lines.append("")

    # Heading hierarchy
    if headings > 0:
        report_lines.append("## Heading Structure")
        for p in doc.paragraphs:
            if p.style.name.startswith("Heading"):
                level = p.style.name.replace("Heading ", "")
                indent = "  " * int(level) if level.isdigit() else ""
                report_lines.append(f"{indent}- [{p.style.name}] {p.text[:80]}")

    if tables > 0:
        report_lines.append("")
        report_lines.append(f"## Table Summary")
        for i, t in enumerate(doc.tables[:5]):
            rows = len(t.rows)
            cols = len(t.columns)
            header_cells = [c.text[:20] for c in t.rows[0].cells] if rows > 0 else []
            report_lines.append(f"- Table {i+1}: {rows} rows × {cols} cols")
            report_lines.append(f"  Headers: {' | '.join(header_cells)}")

    if output_dir:
        os.makedirs(output_dir, exist_ok=True)
        out_path = os.path.join(output_dir, "docx_structure_check.md")
        with open(out_path, "w", encoding="utf-8") as f:
            f.write("\n".join(report_lines))
        print(f"[OK] Check report written to {out_path}")
    else:
        print("\n".join(report_lines))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: check_docx_structure.py <input_docx> [output_dir]")
        sys.exit(1)
    check_structure(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
